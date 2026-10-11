"""Tests for school_register.register and the workbook round trip."""
import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pytest
from conftest import FakeGeocoder, crm, duo_frame, duo_row, geo, pc

from openpyxl import Workbook, load_workbook

from school_register.export import write_dashboard, write_workbook
from school_register.globals import NOTE_ARCHIVED, GlobalSettings
from school_register.matching import Matcher
from school_register.register import (RegisterRow, assign_ids, build_rows, collect_corrections, geocode_records,
                                      read_previous_ids)
from school_register.sources import CbsScore, ExcludedRow, SourceRow, read_cbs_score_table

BESTUREN = {"41572": "De Haagse Scholen", "95694": "Stichting Haagsche Schoolvereeniging"}


def build(tmp_path: Path,
          crm_rows: List[SourceRow],
          pc_rows: List[SourceRow],
          scores: Optional[Dict[str, Dict[str, CbsScore]]] = None) -> Tuple[List[RegisterRow], List[ExcludedRow]]:
    """Run matching, geocoding and field filling on a small DUO set."""
    duo = duo_frame([
        duo_row("05VW00",
                "Haagsche School Vereniging",
                "Nassaulaan",
                "26",
                "2514 JT",
                bg="95694",
                denominatie="Algemeen bijzonder"),
        duo_row("19SI00", "De La Reyschool", "De la Reyweg", "212", "2571 GH"),
    ])
    matcher = Matcher(duo, GlobalSettings())
    matcher.add_crm(crm_rows)
    matcher.add_planned_culture(pc_rows)
    geocoder = FakeGeocoder(tmp_path / "c.json", {
        "2514JT": geo(),
        "2571GH": geo(),
        "2595TG": geo(wijkcode="WK051826"),
        "2280AA": geo(gemeentecode="0603"),
    })
    records, outside = geocode_records(matcher.records, geocoder)
    return build_rows(records, BESTUREN, scores), outside


class TestBuildRows:
    """Tests for build_rows and geocode_records."""

    def test_clean_duo_row_has_no_note(self, tmp_path: Path) -> None:
        """A DUO school confirmed by the CRM has official values, a board name and no note."""
        rows, _ = build(
            tmp_path, [crm("1", "De La Reyschool", "19SI00",
                           ("De la Reyweg", 212, "", "2571GH"), stadsdeel="Centrum")], [])
        row = next(r for r in rows if r.values["VESTIGINGSCODE"] == "19SI00")
        assert row.values["BEVOEGD_GEZAG_NAAM"] == "De Haagse Scholen"
        assert row.values["STADSDEEL"] == "Centrum"
        assert row.values["WIJK"] == "Schildersbuurt"
        assert row.values["DATA_KWALITEIT"] == ""

    def test_nevenlocatie_inherits_from_parent(self, tmp_path: Path) -> None:
        """A sub-location takes board, denominatie and correspondence address from its main location."""
        rows, _ = build(tmp_path, [
            crm("1",
                "Haagsche Schoolvereeniging KSS",
                "05VW00D", ("Koningin Sophiestraat", 24, "A", "2595TG"),
                bg_nummer="95694",
                phone="0703000000"),
            crm("2", "Haagsche Schoolvereeniging NSL", "05VW00B", ("Nassaulaan", 26, "", "2514JT"), bg_nummer="95694"),
        ], [])
        row = next(r for r in rows if r.values["VESTIGINGSCODE"] == "05VW00D")
        assert row.values["BEVOEGD_GEZAG_NUMMER"] == "95694"
        assert row.values["BEVOEGD_GEZAG_NAAM"] == "Stichting Haagsche Schoolvereeniging"
        assert row.values["DENOMINATIE"] == "Algemeen bijzonder"
        assert row.values["STRAATNAAM CORRESPONDENTIEADRES"] == "Postbus"
        assert row.values["TELEFOONNUMMER"] == "0703000000"
        assert row.values["STADSDEEL"] == "Haagse Hout"
        assert row.values["DATA_KWALITEIT"].startswith("Nevenlocatie van 05VW00")

    def test_conflicts_are_noted(self, tmp_path: Path, settings: GlobalSettings) -> None:
        """An unsettled address conflict is a note; a wrong code, archiving and stadsdeel are corrections."""
        rows, _ = build(
            tmp_path, [crm("1", "De La Reyschool", "19SI00", ("De la Reyweg", 212, "", "2571GH"), stadsdeel="Escamp")],
            [pc("2", "De La Reyschool", "19SI01", ("De la Reyweg", 214, "", "2571GH"), archived=True)])
        notes = next(r for r in rows if r.values["VESTIGINGSCODE"] == "19SI00").values["DATA_KWALITEIT"]
        assert "Adres verschilt tussen bronnen" in notes
        assert "Vestigingscode verschilt" not in notes
        assert NOTE_ARCHIVED not in notes
        assert "Stadsdeel" not in notes

        rows, _ = assign_ids(rows, {}, 0, settings)
        fixes = {(c.source, c.field_name): c for c in collect_corrections(rows, [])}
        assert fixes[("Planned Culture", "VESTIGINGSCODE")].value == "19SI01"
        assert fixes[("Planned Culture", "VESTIGINGSCODE")].correct == "19SI00"
        assert fixes[("Planned Culture", "STATUS")].correct == "actief"
        assert fixes[("CRM", "STADSDEEL")].correct == "Centrum"
        assert ("Planned Culture", "ADRES") not in fixes

    def test_outside_municipality_is_excluded(self, tmp_path: Path) -> None:
        """A non-DUO school whose BAG address lies in another municipality is dropped."""
        rows, outside = build(tmp_path, [crm("1", "Elders School", "", ("Laan", 1, "", "2280AA"))], [])
        assert len(rows) == 2
        assert outside[0].reason.startswith("Buiten gemeente 's-Gravenhage")

    def test_missing_location_is_noted(self, tmp_path: Path) -> None:
        """A school whose address and postcode are unknown to the BAG gets a note and no coordinates."""
        rows, _ = build(tmp_path, [crm("1", "Nergens School", "", ("Laan", 1, "", "9999ZZ"))], [])
        row = next(r for r in rows if r.values["VESTIGINGSNAAM"] == "Nergens School")
        assert row.values["LATITUDE"] == ""
        assert "geen kaartlocatie" in row.values["DATA_KWALITEIT"]


class TestAssignIds:
    """Tests for assign_ids and read_previous_ids."""

    def rows(self, tmp_path: Path) -> List[RegisterRow]:
        """Register rows for two DUO schools and one non-DUO school."""
        rows, _ = build(tmp_path, [crm("1", "Lycée Vincent van Gogh", "", ("Scheveningseweg", 237, "", "2514JT"))], [])
        return rows

    def test_first_run(self, tmp_path: Path, settings: GlobalSettings) -> None:
        """A first run numbers DUO schools first, by code, then the others."""
        rows, highest = assign_ids(self.rows(tmp_path), {}, 0, settings)
        assert [r.values["UNIEK_ID"] for r in rows] == ["DH-0001", "DH-0002", "DH-0003"]
        assert [r.values["VESTIGINGSCODE"] for r in rows] == ["05VW00", "19SI00", ""]
        assert highest == 3

    def test_rerun_keeps_ids_and_never_reuses(self, tmp_path: Path, settings: GlobalSettings) -> None:
        """Rerunning through the workbook keeps IDs, and a new school gets a fresh, never reused number."""
        rows, highest = assign_ids(self.rows(tmp_path), {}, 0, settings)
        write_workbook(rows, [], highest + 5, settings)
        previous, stored_highest = read_previous_ids(settings.xlsx_path)
        assert stored_highest == 8

        again, _ = assign_ids(self.rows(tmp_path), previous, stored_highest, settings)
        assert {r.values["VESTIGINGSNAAM"]: r.values["UNIEK_ID"] for r in again} == \
            {r.values["VESTIGINGSNAAM"]: r.values["UNIEK_ID"] for r in rows}

        extra, _ = build(tmp_path, [crm("9", "Nieuwe School", "", ("Laan", 1, "", "2514JT"))], [])
        new = next(r for r in extra if r.values["VESTIGINGSNAAM"] == "Nieuwe School")
        renumbered, _ = assign_ids([new], previous, stored_highest, settings)
        assert renumbered[0].values["UNIEK_ID"] == "DH-0009"

    def test_no_previous_workbook(self, tmp_path: Path) -> None:
        """Without an earlier register there are no IDs to keep."""
        assert read_previous_ids(tmp_path / "missing.xlsx") == ({}, 0)


class TestDashboard:
    """Tests for write_dashboard."""

    def test_data_is_embedded_safely(self, tmp_path: Path, settings: GlobalSettings) -> None:
        """The register is embedded as JSON and a closing script tag in a value cannot break out."""
        rows, _ = assign_ids(TestAssignIds().rows(tmp_path), {}, 0, settings)
        rows[0].values["VESTIGINGSNAAM"] = "School </script><b>x</b>"
        html = write_dashboard(rows, settings).read_text(encoding="utf-8")
        assert "/*__DATA__*/null" not in html
        assert "</script><b>" not in html
        assert re.search(r'const EMBEDDED = \{"columns": \["UNIEK_ID"', html)


class TestCbsScores:
    """Tests for reading CBS achterstandsscores and adding them to register rows."""

    @staticmethod
    def workbook(path: Path) -> Path:
        """Write a small CBS-style workbook with title rows, a header and three schools."""
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Tabel 1"
        sheet.append(["Tabel 1"])
        sheet.append(["Achterstandsscore met drempel per basisschoolvestiging, 1 februari 2025"])
        sheet.append(["School (BRIN-vestiging)", "Aantal bekostigde leerlingen", "Achterstandsscore met drempel"])
        sheet.append(["19SI00", 250, 412.5])
        sheet.append(["05VW00", 38, 0, "a"])
        sheet.append(["99XX00", 120, 12.0, "b, c"])
        workbook.save(path)
        return path

    def test_read_table(self, tmp_path: Path) -> None:
        """Scores are read per vestigingscode; only the "a" marker counts as unreliable."""
        scores = read_cbs_score_table(self.workbook(tmp_path / "cbs.xlsx"))
        assert set(scores) == {"19SI00", "05VW00", "99XX00"}
        assert scores["19SI00"] == CbsScore(score=412.5, unreliable=False)
        assert scores["05VW00"] == CbsScore(score=0.0, unreliable=True)
        assert not scores["99XX00"].unreliable

    def test_missing_sheet(self, tmp_path: Path) -> None:
        """A workbook without table 1 is rejected."""
        path = tmp_path / "other.xlsx"
        Workbook().save(path)
        with pytest.raises(KeyError, match=re.escape("Tabel 1")):
            read_cbs_score_table(path)

    def test_scores_on_duo_rows_only(self, tmp_path: Path) -> None:
        """DUO locations get their score and notes; sub-locations get none, so pupils are not counted twice."""
        table = read_cbs_score_table(self.workbook(tmp_path / "cbs.xlsx"))
        rows, _ = build(tmp_path, [
            crm("1",
                "Haagsche Schoolvereeniging KSS",
                "05VW00D", ("Koningin Sophiestraat", 24, "A", "2595TG"),
                bg_nummer="95694"),
            crm("2", "Haagsche Schoolvereeniging NSL", "05VW00B", ("Nassaulaan", 26, "", "2514JT"), bg_nummer="95694"),
        ], [],
                        scores={
                            "ACHTERSTANDSSCORE_2025": table,
                            "ACHTERSTANDSSCORE_2026_VOORLOPIG": {}
                        })
        by_code = {r.values["VESTIGINGSCODE"]: r.values for r in rows}
        assert by_code["19SI00"]["ACHTERSTANDSSCORE_2025"] == "412.50"
        assert "onbetrouwbaar" not in by_code["19SI00"]["DATA_KWALITEIT"]
        assert by_code["05VW00"]["ACHTERSTANDSSCORE_2025"] == "0.00"
        assert "Achterstandsscore CBS 2025 onbetrouwbaar" in by_code["05VW00"]["DATA_KWALITEIT"]
        assert "Geen CBS-achterstandsscore 2026 (voorlopig)" in by_code["05VW00"]["DATA_KWALITEIT"]
        assert by_code["05VW00D"]["ACHTERSTANDSSCORE_2025"] == ""
        assert "CBS" not in by_code["05VW00D"]["DATA_KWALITEIT"]

    def test_scores_written_as_numbers(self, tmp_path: Path, settings: GlobalSettings) -> None:
        """Score columns are numeric cells in the workbook, so Excel can sort and filter them."""
        table = read_cbs_score_table(self.workbook(tmp_path / "cbs.xlsx"))
        rows, _ = build(tmp_path, [], [], scores={"ACHTERSTANDSSCORE_2025": table})
        rows, highest = assign_ids(rows, {}, 0, settings)
        sheet = load_workbook(write_workbook(rows, [], highest, settings))["Scholen"]
        header = [c.value for c in sheet[1]]
        column = header.index("ACHTERSTANDSSCORE_2025") + 1
        values = [sheet.cell(row=r, column=column).value for r in range(2, sheet.max_row + 1)]
        numbers = [float(v) for v in values if isinstance(v, (int, float))]
        assert sorted(numbers) == [0.0, 412.5]
        assert all(v is None or isinstance(v, (int, float)) for v in values)
