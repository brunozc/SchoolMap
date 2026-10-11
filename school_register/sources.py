"""Loaders for the DUO, CRM and Planned Culture files."""
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pandas as pd
import requests
from openpyxl import load_workbook

from school_register.globals import (CBS_SCORE_FILES, DUO_BESTUREN_URL, DUO_GEMEENTENAAM, SOURCE_CRM, SOURCE_PC,
                                     GlobalSettings)
from school_register.normalize import (Address, clean, is_den_haag, normalize_phone, normalize_postcode,
                                       normalize_website, parse_address_line, parse_split_address, split_code)


@dataclass
class SourceRow:
    """One school as described by the CRM or Planned Culture.

    Attributes:
        source: name of the source file (CRM or Planned Culture).
        source_id: ID of the row in that source.
        name: school name.
        base_code: 6-character DUO-style vestigingscode, empty if none.
        suffix: CRM letter suffix after the base code, empty if none.
        address: main address.
        phone: school phone number, DUO format.
        website: school website, DUO format.
        stadsdeel: stadsdeel as recorded in the source.
        wijk: wijk as recorded in the source.
        bg_nummer: bevoegd gezag number as recorded in the source.
        bg_naam: bevoegd gezag name as recorded in the source.
        archived: True if the source marks the school as archived.
        dependances: extra addresses with their optional vestigingscode (Planned Culture only).
    """
    source: str
    source_id: str
    name: str
    base_code: str
    suffix: str
    address: Address
    phone: str = ""
    website: str = ""
    stadsdeel: str = ""
    wijk: str = ""
    bg_nummer: str = ""
    bg_naam: str = ""
    archived: bool = False
    dependances: List[Tuple[Address, str]] = field(default_factory=list)

    @property
    def full_code(self) -> str:
        """Vestigingscode as written in the source, including the CRM suffix letter."""
        return self.base_code + self.suffix

    def label(self) -> str:
        """Short description for notes and the exclusion list, e.g. "CRM 167 Al - Qoebaschool"."""
        return f"{self.source} {self.source_id} {self.name}"


@dataclass
class ExcludedRow:
    """A source row that is not in the register, with the reason why.

    Attributes:
        source: name of the source file.
        source_id: ID of the row in that source.
        name: school name.
        code: vestigingscode as written in the source.
        address: address as written in the source.
        reason: why the row is excluded.
    """
    source: str
    source_id: str
    name: str
    code: str
    address: str
    reason: str


REASON_TEST = "Testrecord"
REASON_VO = "Voortgezet onderwijs (buiten scope: alleen regulier basisonderwijs)"
REASON_SPECIAL = "Speciaal (basis)onderwijs (buiten scope: alleen regulier basisonderwijs)"
REASON_OUTSIDE = "Buiten gemeente 's-Gravenhage"
REASON_ARCHIVED = "Gearchiveerd in Planned Culture en niet gevonden in CRM of DUO"


def download_file(url: str, path: Path) -> None:
    """Download a file, creating its directory.

    Args:
        url: address of the file.
        path: where to store it.

    Raises:
        requests.HTTPError: if the server returns an error.
    """
    response = requests.get(url, timeout=120)
    response.raise_for_status()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(response.content)


def load_duo(settings: GlobalSettings) -> pd.DataFrame:
    """Load the DUO export and keep the locations in 's-Gravenhage.

    Args:
        settings: build settings.

    Returns:
        The DUO rows for 's-Gravenhage, all columns as strings.
    """
    duo = pd.read_excel(settings.duo_path, dtype=str).fillna("")
    return duo[duo["GEMEENTENAAM"] == DUO_GEMEENTENAAM].reset_index(drop=True)


def load_besturen(settings: GlobalSettings) -> Dict[str, str]:
    """Load the official DUO list of bevoegde gezagen, downloading it on first use.

    Args:
        settings: build settings.

    Returns:
        A mapping from bevoegd gezag number to its official name.

    Raises:
        requests.HTTPError: if the file is not cached and the download fails.
    """
    path = settings.besturen_path
    if not path.exists():
        download_file(DUO_BESTUREN_URL, path)
    besturen = pd.read_csv(path, dtype=str, encoding="utf-8-sig").fillna("")
    return dict(zip(besturen["BEVOEGD GEZAG NUMMER"].map(clean), besturen["BEVOEGD GEZAG NAAM"].map(clean)))


@dataclass(frozen=True)
class CbsScore:
    """Achterstandsscore of one basisschoolvestiging in one CBS table set.

    Attributes:
        score: achterstandsscore met drempel, None if CBS gives no value.
        unreliable: True if CBS marks the score with "a" (40 or fewer pupils).
    """
    score: Optional[float]
    unreliable: bool


def read_cbs_score_table(path: Path) -> Dict[str, CbsScore]:
    """Read table 1 (score met drempel) of a CBS achterstandsscores workbook.

    The read-only reader is used because CBS workbooks can reference images that are missing from the archive, which
    the full reader refuses to open.

    Args:
        path: CBS workbook.

    Returns:
        The score per vestigingscode (BRIN plus location number).

    Raises:
        KeyError: if the workbook has no sheet "Tabel 1".
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Tabel 1"]
    sheet.reset_dimensions()
    scores: Dict[str, CbsScore] = {}
    for row in sheet.iter_rows(values_only=True):
        base, suffix = split_code(row[0] if row else "")
        if not base or suffix:
            continue
        # columns: vestigingscode, number of funded pupils, score met drempel, reliability marker
        raw_score = row[2] if len(row) > 2 else None
        marker = clean(row[3]) if len(row) > 3 else ""
        score = float(raw_score) if isinstance(raw_score, (int, float)) else None
        scores[base] = CbsScore(score=score, unreliable="a" in re.split(r"[,\s]+", marker.lower()))
    workbook.close()
    return scores


def load_cbs_scores(settings: GlobalSettings) -> Dict[str, Dict[str, CbsScore]]:
    """Load every configured CBS achterstandsscores table set, downloading it on first use.

    Args:
        settings: build settings.

    Returns:
        Per output column, the score per vestigingscode.

    Raises:
        requests.HTTPError: if a file is not cached and its download fails.
    """
    tables: Dict[str, Dict[str, CbsScore]] = {}
    for spec in CBS_SCORE_FILES:
        path = settings.cbs_dir / spec.filename
        if not path.exists():
            download_file(spec.url, path)
        tables[spec.column] = read_cbs_score_table(path)
    return tables


def __is_test(name: str, postcode: str) -> bool:
    return bool(re.search(r"test|helpdesk", name, re.IGNORECASE)) or postcode == "1234AB"


def __excluded(source: str, source_id: str, name: str, code: str, address: str, reason: str) -> ExcludedRow:
    """Record an excluded row; test records are masked because their titles often carry colleagues' names."""
    if reason == REASON_TEST:
        return ExcludedRow(source, source_id, "(testrecord)", "", "", reason)
    return ExcludedRow(source, source_id, name, code, address, reason)


def __split_bestuur(raw: object) -> Tuple[str, str]:
    """Split "Stichting X - 40389" into number and name."""
    match = re.match(r"^(.*?)\s*-\s*(\d{5})\s*$", clean(raw))
    if not match:
        return "", ""
    return match.group(2), match.group(1)


def load_crm(settings: GlobalSettings) -> Tuple[List[SourceRow], List[ExcludedRow]]:
    """Load the CRM export.

    Args:
        settings: build settings.

    Returns:
        The schools in scope and the rows that are excluded.
    """
    crm = pd.read_excel(settings.crm_path, dtype=str).fillna("")
    rows: List[SourceRow] = []
    excluded: List[ExcludedRow] = []
    for _, rec in crm.iterrows():
        name = clean(rec["Organisatie"]) or clean(rec["Contact"])
        address = parse_address_line(rec["Adresregel 1"], rec["Postcode"])
        base, suffix = split_code(rec["Vestigingsnummer (BRIN + locatienummer)"])
        reason: Optional[str] = None
        if __is_test(name, address.postcode):
            reason = REASON_TEST
        elif not is_den_haag(rec["Plaats"]):
            reason = REASON_OUTSIDE
        if reason:
            excluded.append(__excluded(SOURCE_CRM, clean(rec["ID"]), name, base + suffix, address.label(), reason))
            continue
        bg_nummer, bg_naam = __split_bestuur(rec["Schoolbestuur - bevoegdgezagnummer"])
        rows.append(
            SourceRow(source=SOURCE_CRM,
                      source_id=clean(rec["ID"]),
                      name=name,
                      base_code=base,
                      suffix=suffix,
                      address=address,
                      phone=normalize_phone(rec["Telefoonnummer"]),
                      website=normalize_website(rec["Website"]),
                      stadsdeel=clean(rec["Stadsdeel"]),
                      wijk=clean(rec["Wijk"]),
                      bg_nummer=bg_nummer,
                      bg_naam=bg_naam))
    return rows, excluded


def __pc_dependances(rec: "pd.Series[str]") -> List[Tuple[Address, str]]:
    dependances: List[Tuple[Address, str]] = []
    for prefix in ("Dependance", "Dependance 2"):
        street = clean(rec[f"{prefix} Straat"])
        if not street:
            continue
        address = parse_split_address(street, rec[f"{prefix} Nummer"], rec[f"{prefix} Toevoeging"],
                                      rec[f"{prefix} Postcode"])
        base, _ = split_code(rec[f"{prefix} Vestigingsnummer"])
        dependances.append((address, base))
    return dependances


def load_planned_culture(settings: GlobalSettings) -> Tuple[List[SourceRow], List[ExcludedRow]]:
    """Load the Planned Culture export and keep regular primary schools in Den Haag.

    The contact-person phone numbers are deliberately not used: they belong to people, not to the school.

    Args:
        settings: build settings.

    Returns:
        The schools in scope and the rows that are excluded.
    """
    pc = pd.read_csv(settings.pc_path, sep=";", dtype=str, encoding="utf-8-sig").fillna("")
    rows: List[SourceRow] = []
    excluded: List[ExcludedRow] = []
    for _, rec in pc.iterrows():
        name = clean(rec["School"])
        address = parse_split_address(rec["Hoofdadres Straat"], rec["Hoofdadres Nummer"], rec["Hoofdadres Toevoeging"],
                                      rec["Hoofdadres Postcode"])
        base, suffix = split_code(rec["Hoofdadres Vestigingsnummer"])
        reason: Optional[str] = None
        if __is_test(name, normalize_postcode(rec["Hoofdadres Postcode"])):
            reason = REASON_TEST
        elif clean(rec["Onderwijstype"]) != "Primair Onderwijs":
            reason = REASON_VO
        elif clean(rec["Speciaal onderwijs"]) == "Ja":
            reason = REASON_SPECIAL
        elif not is_den_haag(rec["Hoofdadres Plaats"]):
            reason = REASON_OUTSIDE
        if reason:
            excluded.append(__excluded(SOURCE_PC, clean(rec["Id"]), name, base + suffix, address.label(), reason))
            continue
        rows.append(
            SourceRow(source=SOURCE_PC,
                      source_id=clean(rec["Id"]),
                      name=name,
                      base_code=base,
                      suffix=suffix,
                      address=address,
                      bg_naam=clean(rec["Schoolbestuur"]),
                      archived=clean(rec["Gearchiveerd"]) == "Ja",
                      dependances=__pc_dependances(rec)))
    return rows, excluded
