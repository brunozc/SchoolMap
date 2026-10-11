"""Turn matched school records into register rows: field precedence, DATA_KWALITEIT notes and UNIEK_ID."""
import re
from dataclasses import dataclass, field
from itertools import combinations
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

from openpyxl import load_workbook

from school_register.geocode import PdokGeocoder
from school_register.globals import (CBS_SCORE_FILES, DUO_COPIED_COLUMNS, DUO_GEMEENTENAAM, DUO_PLAATSNAAM,
                                     DUO_TO_OUTPUT, GEMEENTE_CODE, INHERITED_COLUMNS, KIND_DUO, NOTE_ADDRESS_CONFLICT,
                                     NOTE_ARCHIVED, NOTE_BG_NOT_FROM_DUO, NOTE_CODE_CONFLICT, NOTE_NO_LOCATION,
                                     NOTE_NOT_IN_BAG, NOTE_ONLY_PC, NOTE_PHONE_INVALID, NOTE_SCORE_MISSING,
                                     NOTE_SCORE_UNRELIABLE, OUTPUT_COLUMNS, SOURCE_CRM, SOURCE_DUO, SOURCE_PC,
                                     SOURCE_RIO, WIJK_STADSDEEL, GlobalSettings)
from school_register.matching import SchoolRecord
from school_register.normalize import (Address, clean, format_house_number, format_postcode, is_valid_phone,
                                       normalize_postcode, normalized_name, parse_house_number)
from school_register.rio import RioBuilding
from school_register.sources import REASON_OUTSIDE, CbsScore, ExcludedRow

ID_SHEET = "Scholen"
LINK_SHEET = "Koppeling"
INFO_SHEET = "Toelichting"
HIGHEST_ID_LABEL = "Hoogste toegekende UNIEK_ID"


@dataclass(frozen=True)
class Correction:
    """A value in the CRM or Planned Culture that disagrees with the official sources.

    Attributes:
        source: source file with the wrong value (or RIO for a building missing from the files).
        source_id: ID of the row in that source.
        source_name: school name in that source.
        uid: UNIEK_ID of the register row.
        field_name: which value is wrong.
        value: the value in the source.
        correct: the official value.
        authority: where the official value comes from.
    """
    source: str
    source_id: str
    source_name: str
    uid: str
    field_name: str
    value: str
    correct: str
    authority: str


@dataclass
class RegisterRow:
    """A finished row of the register.

    Attributes:
        record: the matched school record behind this row.
        values: output values per column, including LATITUDE and LONGITUDE.
    """
    record: SchoolRecord
    values: Dict[str, str] = field(default_factory=dict)


def geocode_records(records: List[SchoolRecord],
                    geocoder: PdokGeocoder) -> Tuple[List[SchoolRecord], List[ExcludedRow]]:
    """Look up every record in the BAG and drop non-DUO records that lie outside the municipality.

    Args:
        records: matched school records.
        geocoder: PDOK geocoder.

    Returns:
        The records inside the municipality and the excluded source rows of those outside it.
    """
    kept: List[SchoolRecord] = []
    excluded: List[ExcludedRow] = []
    for record in records:
        record.geo = geocoder.geocode(record.address)
        if record.kind != KIND_DUO and record.geo and record.geo.exact and record.geo.gemeentecode != GEMEENTE_CODE:
            reason = f"{REASON_OUTSIDE} (BAG: {record.geo.gemeentenaam})"
            rows = record.sources or ([record.dependance_of] if record.dependance_of else [])
            for row in rows:
                excluded.append(
                    ExcludedRow(row.source, row.source_id, row.name, row.full_code, row.address.label(), reason))
            continue
        kept.append(record)
    geocoder.save()
    return kept, excluded


def __first(record: SchoolRecord, attribute: str, source: Optional[str] = None) -> str:
    """First non-empty value of a SourceRow attribute, CRM rows before Planned Culture rows."""
    rows = sorted(record.sources, key=lambda r: r.source != SOURCE_CRM)
    for row in rows:
        if source and row.source != source:
            continue
        value = getattr(row, attribute)
        if value:
            return str(value)
    return ""


def __duo_values(record: SchoolRecord) -> Dict[str, str]:
    return {DUO_TO_OUTPUT.get(col, col): record.duo.get(col, "") for col in DUO_COPIED_COLUMNS}


def __own_values(record: SchoolRecord, parent_values: Optional[Dict[str, str]]) -> Dict[str, str]:
    """Values of a location that is not in DUO, from its sources, the BAG and its main location."""
    values = {col: "" for col in OUTPUT_COLUMNS}
    geo = record.geo
    if geo and geo.exact:
        street, number, addition, postcode = geo.street, geo.number, geo.addition, geo.postcode
    else:
        street, number = record.address.street, record.address.number
        addition, postcode = record.address.addition, record.address.postcode
    sourced_name = __first(record, "name", SOURCE_CRM) if record.sources else ""
    values.update({
        "VESTIGINGSNAAM": sourced_name or record.name,
        "VESTIGINGSCODE": record.code,
        "INSTELLINGSCODE": record.code[:4],
        "STRAATNAAM": street,
        "HUISNUMMER-TOEVOEGING": format_house_number(number, addition),
        "POSTCODE": format_postcode(postcode),
        "PLAATSNAAM": DUO_PLAATSNAAM,
        "GEMEENTENUMMER": GEMEENTE_CODE,
        "GEMEENTENAAM": DUO_GEMEENTENAAM,
        "TELEFOONNUMMER": __first(record, "phone"),
        "INTERNETADRES": __first(record, "website"),
    })
    if values["TELEFOONNUMMER"] and not is_valid_phone(values["TELEFOONNUMMER"]):
        record.note(NOTE_PHONE_INVALID)
    if record.inherit and parent_values is not None:
        values.update({col: parent_values[col] for col in INHERITED_COLUMNS})
    elif record.rio_bestuur[0]:
        values["BEVOEGD_GEZAG_NUMMER"], values["BEVOEGD_GEZAG_NAAM"] = record.rio_bestuur
    else:
        values["BEVOEGD_GEZAG_NUMMER"] = __first(record, "bg_nummer", SOURCE_CRM)
    if record.rio_name:
        values["VESTIGINGSNAAM"] = record.rio_name
    return values


def __apply_bestuur_name(record: SchoolRecord, values: Dict[str, str], besturen: Dict[str, str]) -> None:
    nummer = values["BEVOEGD_GEZAG_NUMMER"]
    if not nummer or values["BEVOEGD_GEZAG_NAAM"]:
        return
    if nummer in besturen:
        values["BEVOEGD_GEZAG_NAAM"] = besturen[nummer]
        return
    crm_name = next((r.bg_naam for r in record.sources if r.source == SOURCE_CRM and r.bg_nummer == nummer), "")
    if crm_name:
        values["BEVOEGD_GEZAG_NAAM"] = crm_name
        record.note(NOTE_BG_NOT_FROM_DUO)


def __apply_area(record: SchoolRecord, values: Dict[str, str]) -> None:
    geo = record.geo
    values["WIJK"] = geo.wijknaam if geo else ""
    derived = WIJK_STADSDEEL.get(geo.wijkcode, "") if geo else ""
    # a CRM stadsdeel that disagrees with the BAG wijk is reported in the corrections, the wijk wins
    values["STADSDEEL"] = derived or __first(record, "stadsdeel", SOURCE_CRM)
    point: Optional[Tuple[float, float]] = (geo.lat, geo.lon) if geo else None
    if geo is not None and geo.exact:
        pass
    elif record.rio_point is not None:
        # RIO confirms this very address and has its BAG location, PDOK only found the postcode
        point = record.rio_point
    elif geo is None:
        record.note(NOTE_NO_LOCATION)
    else:
        record.note(NOTE_NOT_IN_BAG)
    values["LATITUDE"] = f"{point[0]:.6f}" if point else ""
    values["LONGITUDE"] = f"{point[1]:.6f}" if point else ""


def __apply_conflicts(record: SchoolRecord) -> None:
    """Note disagreements about code and address that no official source settles, and archived or PC-only schools.

    When DUO or RIO confirms the code or address, the disagreeing source rows are reported as corrections instead.
    """
    codes: List[Tuple[str, str]] = []
    addresses: List[Tuple[str, Address]] = []
    if record.kind == KIND_DUO:
        codes.append((SOURCE_DUO, record.code))
        addresses.append((SOURCE_DUO, record.address))
    for row in record.sources:
        if row.base_code:
            codes.append((row.source, row.full_code))
        if row.address.key:
            addresses.append((row.source, row.address))
    if not record.code_official and len({code[:6] for _, code in codes}) > 1:
        record.note(NOTE_CODE_CONFLICT.format(codes=" / ".join(f"{s} {c}" for s, c in dict.fromkeys(codes))))
    if not record.address_confirmed and any(not a.same_place(b) for (_, a), (_, b) in combinations(addresses, 2)):
        unique = dict.fromkeys((s, a.label()) for s, a in addresses)
        record.note(NOTE_ADDRESS_CONFLICT.format(addresses=" / ".join(f"{s} {label}" for s, label in unique)))
    active = record.kind == KIND_DUO or record.address_confirmed or record.code_official
    if not active and any(r.archived for r in record.sources if r.source == SOURCE_PC):
        record.note(NOTE_ARCHIVED)
    if not record.address_confirmed and record.kind != KIND_DUO and record.sources and not record.has_source(
            SOURCE_CRM):
        record.note(NOTE_ONLY_PC)


def __apply_scores(record: SchoolRecord, values: Dict[str, str], scores: Dict[str, Dict[str, CbsScore]]) -> None:
    """Fill the CBS achterstandsscores; only DUO locations have one, so sub-locations are not counted twice."""
    for spec in CBS_SCORE_FILES:
        values[spec.column] = ""
        if record.kind != KIND_DUO or spec.column not in scores:
            continue
        found = scores[spec.column].get(record.code)
        if found is None or found.score is None:
            record.note(NOTE_SCORE_MISSING.format(year=spec.label))
            continue
        values[spec.column] = f"{found.score:.2f}"
        if found.unreliable:
            record.note(NOTE_SCORE_UNRELIABLE.format(year=spec.label))


def build_rows(records: List[SchoolRecord],
               besturen: Dict[str, str],
               scores: Optional[Dict[str, Dict[str, CbsScore]]] = None) -> List[RegisterRow]:
    """Fill the output columns of every record and collect its DATA_KWALITEIT notes.

    DUO values win for every DUO column. Locations outside DUO take name and contact details from the CRM (then
    Planned Culture), their address from the BAG, and board, denominatie and correspondence address from their main
    location when they are a sub-location or dependance. CBS achterstandsscores are added to DUO locations only.

    Args:
        records: geocoded school records.
        besturen: official names of bevoegde gezagen by number.
        scores: CBS achterstandsscores per output column and vestigingscode; None leaves the score columns empty.

    Returns:
        One register row per record, in the order of the records; UNIEK_ID is still empty.
    """
    done: Dict[int, Dict[str, str]] = {}

    def values_of(record: SchoolRecord) -> Dict[str, str]:
        if id(record) in done:
            return done[id(record)]
        if record.kind == KIND_DUO:
            values = {col: "" for col in OUTPUT_COLUMNS}
            values.update(__duo_values(record))
        else:
            parent_values = values_of(record.parent) if record.parent is not None else None
            values = __own_values(record, parent_values)
        __apply_bestuur_name(record, values, besturen)
        done[id(record)] = values
        return values

    rows: List[RegisterRow] = []
    for record in records:
        values = dict(values_of(record))
        __apply_conflicts(record)
        __apply_area(record, values)
        __apply_scores(record, values, scores or {})
        values["DATA_KWALITEIT"] = "; ".join(record.notes)
        rows.append(RegisterRow(record=record, values=values))
    return rows


def collect_corrections(rows: List[RegisterRow], unlisted: List[Tuple[SchoolRecord, RioBuilding]]) -> List[Correction]:
    """List the values to fix in the CRM and Planned Culture, now that DUO, RIO and the BAG settle them.

    Only values with an official answer are listed: codes confirmed by DUO or RIO, addresses that RIO registers as a
    building of the school, stadsdelen from the BAG wijk, invalid phone numbers, schools archived in Planned Culture
    that DUO or RIO shows as active, and RIO buildings that no register row covers.

    Args:
        rows: register rows with UNIEK_ID.
        unlisted: RIO buildings of register schools without a register row, with their school.

    Returns:
        The corrections, ordered by UNIEK_ID.
    """
    corrections: List[Correction] = []
    for row in rows:
        record, uid = row.record, row.values["UNIEK_ID"]
        authority = SOURCE_DUO if record.kind == KIND_DUO else SOURCE_RIO
        active = record.kind == KIND_DUO or record.address_confirmed or record.code_official
        for src in record.sources:

            def add(field_name: str, value: str, correct: str, by: str) -> None:
                corrections.append(Correction(src.source, src.source_id, src.name, uid, field_name, value, correct, by))

            if record.code_official and src.base_code != record.code[:6]:
                add("VESTIGINGSCODE", src.full_code or "(leeg)", record.code, authority)
            if record.address_confirmed and src.address.key and not src.address.same_place(record.address):
                add("ADRES", src.address.label(), record.address.label(),
                    f"{authority} + RIO-gebouw" if authority == SOURCE_DUO else "RIO-gebouw")
            if src.stadsdeel and row.values["STADSDEEL"] and src.stadsdeel != row.values["STADSDEEL"]:
                add("STADSDEEL", src.stadsdeel, row.values["STADSDEEL"], f"BAG-wijk {row.values['WIJK']}")
            if src.phone and not is_valid_phone(src.phone):
                add("TELEFOONNUMMER", src.phone, "controleren: geen geldig Nederlands nummer", "formaat")
            if src.archived and active:
                add("STATUS", "gearchiveerd", "actief", authority)
    uid_of = {id(r.record): r.values["UNIEK_ID"] for r in rows}
    for owner, building in unlisted:
        corrections.append(
            Correction(SOURCE_RIO, building.code, owner.rio_name or owner.name,
                       uid_of.get(id(owner), ""), "ONTBREKEND GEBOUW", "(niet in CRM, Planned Culture of register)",
                       building.address.label(), "RIO-gebouw van deze school"))
    return sorted(corrections, key=lambda c: (__id_number(c.uid), c.source, c.field_name))


def __source_keys(record: SchoolRecord, name: str) -> List[Tuple[str, ...]]:
    """Keys from the source rows behind a register row, the most stable way to recognise it on a later run."""
    keys: List[Tuple[str, ...]] = []
    if record.kind == KIND_DUO:
        keys.append(("source", SOURCE_DUO, record.code))
    keys += [("source", s.source, s.source_id) for s in record.sources]
    if record.dependance_of is not None:
        keys.append(("source", f"{record.dependance_of.source} (dependance)", record.dependance_of.source_id,
                     normalized_name(name)))
    return keys


def __id_keys(values: Dict[str, str]) -> List[Tuple[str, ...]]:
    """Keys under which a row is recognised again on a later run."""
    postcode = normalize_postcode(values.get("POSTCODE", ""))
    number, _ = parse_house_number(values.get("HUISNUMMER-TOEVOEGING", ""))
    keys: List[Tuple[str, ...]] = []
    code = clean(values.get("VESTIGINGSCODE", ""))
    if code:
        keys.append(("code", code, postcode, str(number)))
    keys.append(("name", normalized_name(clean(values.get("VESTIGINGSNAAM", ""))), postcode))
    return keys


def read_previous_ids(path: Path) -> Tuple[Dict[Tuple[str, ...], str], int]:
    """Read the UNIEK_IDs of an earlier register, so that reruns keep them.

    Args:
        path: earlier output workbook; it may not exist.

    Returns:
        The IDs by matching key, and the highest ID number ever assigned (0 if there is no earlier register).
    """
    if not path.exists():
        return {}, 0
    workbook = load_workbook(path, read_only=True, data_only=True)
    by_key: Dict[Tuple[str, ...], str] = {}
    highest = 0
    sheet = workbook[ID_SHEET]
    rows = sheet.iter_rows(values_only=True)
    header = [clean(h) for h in next(rows)]
    for raw in rows:
        values = {h: clean(v) for h, v in zip(header, raw)}
        uid = values.get("UNIEK_ID", "")
        if not uid:
            continue
        for key in __id_keys(values):
            by_key.setdefault(key, uid)
        highest = max(highest, __id_number(uid))
    if LINK_SHEET in workbook.sheetnames:
        links = workbook[LINK_SHEET].iter_rows(values_only=True)
        link_header = [clean(h) for h in next(links)]
        for raw in links:
            link = {h: clean(v) for h, v in zip(link_header, raw)}
            source_key: Tuple[str, ...] = ("source", link["BRON"], link["BRON_ID"])
            if link["BRON"].endswith("(dependance)"):
                source_key = source_key + (normalized_name(link["VESTIGINGSNAAM"]), )
            if link["UNIEK_ID"]:
                by_key.setdefault(source_key, link["UNIEK_ID"])
    if INFO_SHEET in workbook.sheetnames:
        for label, value, *_ in workbook[INFO_SHEET].iter_rows(values_only=True):
            if clean(label) == HIGHEST_ID_LABEL:
                highest = max(highest, __id_number(clean(value)))
    workbook.close()
    return by_key, highest


def __id_number(uid: str) -> int:
    match = re.search(r"(\d+)$", uid)
    return int(match.group(1)) if match else 0


def assign_ids(rows: List[RegisterRow], previous: Dict[Tuple[str, ...], str], highest: int,
               settings: GlobalSettings) -> Tuple[List[RegisterRow], int]:
    """Give every row a UNIEK_ID, keeping the IDs of an earlier run and never reusing a number.

    A row keeps its earlier ID when one of its source rows (DUO code, CRM ID, Planned Culture ID) was linked to it
    before, or else when code and address, or name and postcode, match. New rows are numbered DUO locations first (by
    vestigingscode), then the others by name.

    Args:
        rows: register rows without UNIEK_ID.
        previous: IDs of an earlier run by matching key.
        highest: highest ID number ever assigned.
        settings: build settings.

    Returns:
        The rows sorted by UNIEK_ID, and the new highest ID number.
    """
    ordered = sorted(rows,
                     key=lambda r: (r.record.kind != KIND_DUO, r.values["VESTIGINGSCODE"]
                                    if r.record.kind == KIND_DUO else "", r.values["VESTIGINGSNAAM"].lower()))
    used: Set[str] = set()
    for row in ordered:
        keys = __source_keys(row.record, row.values["VESTIGINGSNAAM"]) + __id_keys(row.values)
        uid = next((previous[k] for k in keys if k in previous and previous[k] not in used), "")
        if uid:
            row.values["UNIEK_ID"] = uid
            used.add(uid)
    for row in ordered:
        if not row.values["UNIEK_ID"]:
            highest += 1
            row.values["UNIEK_ID"] = f"{settings.id_prefix}-{highest:0{settings.id_digits}d}"
    return sorted(ordered, key=lambda r: __id_number(r.values["UNIEK_ID"])), highest
